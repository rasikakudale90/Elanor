import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_strict_assignment2_document(filename="Assignment_2_User_Research_and_Competitive_Analysis.docx"):
    doc = docx.Document()

    # Colors
    C_PRIMARY = RGBColor(27, 26, 23)      # #1B1A17 - Dark Charcoal
    C_SECONDARY = RGBColor(70, 85, 65)    # #465541 - Deep Sage
    C_ACCENT = RGBColor(160, 110, 40)     # #A06E28 - Warm Bronze
    C_BODY = RGBColor(40, 40, 40)         # Clean dark readable text

    # Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Table formatting helpers
    def set_cell_background(cell, fill_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_padding(cell, top=75, bottom=75, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    # Headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.color.rgb = C_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.color.rgb = C_SECONDARY
        return p

    def add_body(text, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = C_BODY
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.name = 'Calibri'
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = C_PRIMARY

        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = C_BODY
        return p

    def create_table(headers, rows_data, widths=None):
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Header Row
        hdr = tbl.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = h
            set_cell_background(hdr[i], "1B1A17")
            set_cell_padding(hdr[i], top=70, bottom=70, left=90, right=90)
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_i, row in enumerate(rows_data):
            cells = tbl.rows[r_i + 1].cells
            bg = "FAF7F2" if r_i % 2 == 1 else "FFFFFF"
            for c_i, val in enumerate(row):
                cells[c_i].text = str(val)
                set_cell_background(cells[c_i], bg)
                set_cell_padding(cells[c_i], top=60, bottom=60, left=90, right=90)
                p = cells[c_i].paragraphs[0]
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9)
                    r.font.color.rgb = C_BODY

        if widths:
            for row in tbl.rows:
                for c_i, w in enumerate(widths):
                    row.cells[c_i].width = Inches(w)

        tblPr = tbl._tbl.tblPr
        b = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D3C5B4"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3C5B4"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="EFE7DD"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(b)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return tbl

    # ==========================================
    # ACADEMIC HEADER
    # ==========================================
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(2)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sub = p_meta.add_run("MCA DEGREE PROGRAM | UI/UX DESIGN")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(8.5)
    r_sub.bold = True
    r_sub.font.color.rgb = C_ACCENT

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_t1 = p_title.add_run("ASSIGNMENT 2: USER RESEARCH & COMPETITIVE ANALYSIS\n")
    r_t1.bold = True
    r_t1.font.name = 'Arial'
    r_t1.font.size = Pt(15)
    r_t1.font.color.rgb = C_PRIMARY

    r_t2 = p_title.add_run("Application Case Study: Élanor (Luxury Skincare & AI Skin Diagnostic)")
    r_t2.italic = True
    r_t2.font.name = 'Georgia'
    r_t2.font.size = Pt(10.5)
    r_t2.font.color.rgb = C_SECONDARY

    # Metadata Block
    info_headers = ["Field", "Submission Details"]
    info_rows = [
        ["Course", "Master of Computer Applications (MCA)"],
        ["Subject", "UI/UX Design"],
        ["Student Name", "Rasika Kudale"],
        ["Roll Number / PRN", "[Your Roll Number / PRN]"],
        ["Assignment Scope", "User Research (5 Users), Competitive Analysis (2 Apps), Personas, Empathy Map, Journey Map, Problem Statement, HMW Statements, User Stories & Acceptance Criteria"]
    ]
    create_table(info_headers, info_rows, widths=[2.0, 4.5])

    # ==========================================
    # 1. USER RESEARCH WITH FIVE USERS
    # ==========================================
    add_h1("1. User Research with Five Users")
    add_body(
        "I conducted qualitative user research with 5 diverse participants to understand their skincare buying habits, how they navigate skincare websites on mobile/desktop, and what issues they encounter during online shopping."
    )

    u_headers = ["User", "Profile & Age", "Skincare Focus & Habits", "Primary Device", "Key Feedback & Pain Point"]
    u_rows = [
        ["User 1", "Genevieve (34)\nSenior Director", "Buys premium anti-aging serums monthly; very busy schedule.", "iPhone (Mobile Safari)", "Wants checkout in under 2 minutes. Hates reaching for top-corner menus with one hand."],
        ["User 2", "Claire (26)\nUI Designer", "Reads ingredient labels (Niacinamide, Squalane) and active percentages.", "Laptop & Mobile", "Loves the AI Skin Concierge laser scan. Wants a visible step-by-step progress bar during the quiz."],
        ["User 3", "Ananya (29)\nProduct Manager", "Sensitive combination skin; seeks clear product guidance.", "Mobile Chrome", "Confused when websites split products into 'By Category' and 'By Concern' without cross-links."],
        ["User 4", "Marcus (38)\nArchitect", "Simple routine (cleanser + moisturizer); values clean aesthetic.", "iPad & Mobile", "Abandons carts if clicking 'Add to Cart' gives no instant confirmation or if shipping charges are hidden."],
        ["User 5", "Pooja (23)\nStudent", "Budget-conscious; actively looks for promo codes and free delivery.", "Mobile Safari", "Wants to see how close she is to the Free Shipping threshold directly inside the cart drawer."]
    ]
    create_table(u_headers, u_rows, widths=[0.8, 1.3, 1.4, 1.2, 1.8])

    add_h2("1.1 Summary of Core Research Insights")
    add_bullet("Mobile Ergonomics Matter: ", "3 out of 5 users shopped with one hand on mobile and found top headers hard to reach, validating the need for a bottom navigation bar.")
    add_bullet("Interactive Diagnostics Build Trust: ", "All users preferred the 3D-parallax AI Skin Concierge over boring text surveys because it felt like a personalized clinic consultation.")
    add_bullet("Transparent Cart & Shipping Incentives: ", "Users strongly responded to free shipping progress bars and 1-click coupon pills inside the cart drawer.")

    # ==========================================
    # 2. COMPETITIVE ANALYSIS (TWO COMPETING APPLICATIONS)
    # ==========================================
    add_h1("2. Competitive Analysis: Comparing Two Competing Applications")
    add_body(
        "I selected and compared two leading applications in the luxury and clinical skincare space against Élanor: Aesop (luxury minimalist brand) and Kiehl's (clinical heritage apothecary brand)."
    )

    c_headers = ["Feature / UX Aspect", "Competitor 1: Aesop", "Competitor 2: Kiehl's", "Élanor (Our Platform)"]
    c_rows = [
        ["Visual Aesthetic", "Minimalist, amber tones, editorial photography.", "Apothecary style, clinical badges, high visual clutter.", "Quiet luxury, alabaster linen palette, Cormorant Garamond headings."],
        ["Navigation", "Unconventional nested menus; high learning curve.", "Traditional mega-menu; heavy on mobile viewports.", "Simple top header + Persistent Mobile Bottom Bar (`MobileBottomNav.tsx`)."],
        ["Skin Diagnostic Tool", "Static text recommendations; no interactive tool.", "Standard multi-choice survey; plain product links.", "Interactive 3D model portrait, laser scanning animation, Dermal Vitality Matrix."],
        ["Ingredient Transparency", "Lists botanical notes, but lacks active percentages.", "Good ingredient list, but lacks holistic synergy view.", "Botanical active synergy accordion with clinical percentage proof."],
        ["Cart Feedback", "Full-page redirect; no instant sliding drawer.", "Slide-out drawer with aggressive popups.", "Slide-out Cart Drawer (`CartDrawer.tsx`) with Free Shipping progress & coupon pills."],
        ["Checkout Flow", "Multi-step checkout with multiple page reloads.", "3-step checkout requiring account sign-in.", "Single-Page Accordion Checkout with 1-click VIP test profiles and sandbox payments."]
    ]
    create_table(c_headers, c_rows, widths=[1.5, 1.7, 1.7, 1.6])

    add_h2("2.1 Competitive Takeaway")
    add_body(
        "Aesop has great luxury aesthetics but lacks interactive diagnostic technology. Kiehl's has clinical skincare tools but has high visual clutter. Élanor combines the clean quiet luxury of Aesop with the scientific diagnostic power of Kiehl's, delivered through an easy, modern interface."
    )

    # ==========================================
    # 3. USER PERSONAS
    # ==========================================
    add_h1("3. User Personas")
    add_body(
        "Based on the user research findings, I created two user personas representing the primary customer segments for Élanor:"
    )

    add_h2("3.1 Persona 1: Genevieve Moreau (The Busy Luxury Shopper)")
    add_bullet("Age & Role: ", "34 years old, Senior Corporate Director.")
    add_bullet("Goals: ", "Wants clean, premium cellular skincare for anti-aging. Needs a frictionless 2-minute mobile shopping experience during her daily commute.")
    add_bullet("Frustrations: ", "Hates cluttered websites, hidden shipping costs, and complicated multi-step checkout forms.")
    add_bullet("Preferred Device: ", "iPhone (uses one-handed browsing 85% of the time).")

    add_h2("3.2 Persona 2: Claire Delacroix (The Ingredient & AI Enthusiast)")
    add_bullet("Age & Role: ", "26 years old, Digital Product Designer.")
    add_bullet("Goals: ", "Wants to understand active ingredients (Niacinamide, Squalane) and take the AI Skin Concierge quiz to receive a tailored 7-day AM/PM routine.")
    add_bullet("Frustrations: ", "Skeptical of generic marketing claims; dislikes quizzes that don't show how many questions are left.")
    add_bullet("Preferred Device: ", "MacBook Pro and Android smartphone.")

    # ==========================================
    # 4. EMPATHY MAP
    # ==========================================
    add_h1("4. Empathy Map")
    add_body(
        "I mapped the target user's thoughts, feelings, words, and actions into a 4-quadrant Empathy Map:"
    )

    emp_headers = ["Quadrant", "User Realities & Observations"]
    emp_rows = [
        [
            "SAYS",
            "• 'I want to see exactly what active botanicals are in this formula.'\n"
            "• 'If the checkout takes more than 2 minutes on my phone, I will abandon it.'\n"
            "• 'I want an AI skin diagnostic that feels personalized, not just a generic survey.'\n"
            "• 'Tell me upfront how much more I need to buy to get free shipping.'"
        ],
        [
            "THINKS",
            "• 'Is this luxury skincare backed by real science or just fancy branding?'\n"
            "• 'Will these ingredients irritate my sensitive skin barrier?'\n"
            "• 'My evening skincare routine should feel calm, luxurious, and easy.'\n"
            "• 'I hope there are no hidden delivery charges added at the last second.'"
        ],
        [
            "DOES",
            "• Browses products on her phone during commutes or evening downtime.\n"
            "• Clicks into ingredient accordions to verify percentages and clinical results.\n"
            "• Uses the interactive AI Skin Concierge quiz to get tailored recommendations.\n"
            "• Taps 1-click discount coupon pills inside the cart drawer."
        ],
        [
            "FEELS",
            "• Overwhelmed when websites have too many products without guidance.\n"
            "• Anxious about skin aging, hydration loss, and skin sensitivity.\n"
            "• Delighted when a website is fast, beautifully designed, and easy to navigate.\n"
            "• Confident when clinical percentages and customer reviews validate a product."
        ]
    ]
    create_table(emp_headers, emp_rows, widths=[1.5, 5.0])

    # ==========================================
    # 5. CUSTOMER JOURNEY MAP
    # ==========================================
    doc.add_page_break()
    add_h1("5. Customer Journey Map")
    add_body(
        "This journey map tracks the user's end-to-end experience across 5 stages of interacting with Élanor:"
    )

    cjm_headers = ["Stage", "User Action", "Touchpoint", "User Emotion", "Pain Point Found", "UX Solution"]
    cjm_rows = [
        ["1. Discovery", "Lands on home screen; browses brand story.", "Homepage (`/`)", "Interested 😊", "Top menu hard to reach with one thumb.", "Add Mobile Bottom Nav bar."],
        ["2. Consultation", "Takes AI quiz for skin recommendation.", "AI Concierge (`/ai-skin-concierge`)", "Engaged 😊", "Quiz feels long if progress is hidden.", "Add 4-step progress stepper."],
        ["3. Evaluation", "Checks product details & ingredients.", "Product Page (`/product/[id]`)", "Informed 😐", "Small grey text hard to read.", "Darken text to pass WCAG AA."],
        ["4. Purchase", "Adds item to bag, applies coupon, pays.", "Cart Drawer & Checkout (`/checkout`)", "Decisive 😊", "Silent click; long multi-page checkout.", "Slide-out drawer + 1-page checkout."],
        ["5. Tracking", "Checks package delivery status.", "Order Tracker (`/track`)", "Satisfied 😊", "Uncertainty about shipping milestone.", "5-stage visual milestone timeline."]
    ]
    create_table(cjm_headers, cjm_rows, widths=[1.0, 1.2, 1.2, 0.9, 1.1, 1.1])

    # ==========================================
    # 6. PROBLEM STATEMENT
    # ==========================================
    add_h1("6. Problem Statement")
    add_body(
        "Busy and skin-conscious consumers need an intuitive luxury skincare web platform that provides transparent clinical diagnostics and fast mobile checkout, but currently they face low-contrast unreadable text, awkward mobile navigation, silent cart feedback, and long multi-step checkouts that cause frustration and cart abandonment."
    )

    # ==========================================
    # 7. HOW MIGHT WE (HMW) STATEMENTS
    # ==========================================
    add_h1("7. How Might We (HMW) Statements")
    add_body(
        "Based on the research findings and problem statement, I formulated 5 actionable 'How Might We' design questions:"
    )
    add_bullet("HMW 1: ", "How might we make an online skin consultation feel as personal and engaging as visiting a dermatologist?")
    add_bullet("HMW 2: ", "How might we make mobile navigation comfortable for users holding large phones with one hand?")
    add_bullet("HMW 3: ", "How might we display clinical ingredients and percentages in a clear, high-contrast, easy-to-read format?")
    add_bullet("HMW 4: ", "How might we give users immediate visual confirmation when they add an item to their bag while encouraging higher cart value?")
    add_bullet("HMW 5: ", "How might we compress the checkout process into a fast single-page flow that completes in under 60 seconds?")

    # ==========================================
    # 8. USER STORIES AND ACCEPTANCE CRITERIA
    # ==========================================
    add_h1("8. User Stories and Acceptance Criteria")
    add_body(
        "Here are 5 Agile User Stories with testable Acceptance Criteria that directly answer our HMW questions:"
    )

    # User Story 1
    add_h2("User Story 1: AI Skin Consultation Quiz")
    add_body("As a skin-conscious user (Claire),\nI want to complete an interactive 4-step skin diagnostic quiz,\nSo that I can receive an objective, personalized skincare routine for my skin type.")
    add_bullet("Acceptance Criteria 1.1: ", "When the user opens `/ai-skin-concierge`, the page must show an interactive model portrait with a scanning laser animation.")
    add_bullet("Acceptance Criteria 1.2: ", "The quiz must display a clear step indicator ('Step X of 4') and a live summary box that records selections in real time.")
    add_bullet("Acceptance Criteria 1.3: ", "Upon completion, the app must display Dermal Vitality Matrix scores and recommend a personalized Morning/Evening routine.")
    add_bullet("Acceptance Criteria 1.4: ", "Clicking 'Add Entire Ritual to Bag' must add all recommended items to the cart with a 15% discount applied.")

    # User Story 2
    add_h2("User Story 2: Mobile Bottom Navigation Bar")
    add_body("As a mobile shopper (Genevieve),\nI want a persistent bottom navigation bar on my phone,\nSo that I can easily tap Home, Shop, Concierge, and Cart with my thumb without stretching.")
    add_bullet("Acceptance Criteria 2.1: ", "On mobile screens (<1024px), a bottom bar with Home, Shop, Concierge, and Profile must remain fixed at the screen bottom.")
    add_bullet("Acceptance Criteria 2.2: ", "A floating Bag button on the bar must show the live item count badge whenever items are in the cart.")
    add_bullet("Acceptance Criteria 2.3: ", "Tapping any navigation icon must switch pages instantly without lag.")

    # User Story 3
    add_h2("User Story 3: Slide-Out Cart Drawer with Free Shipping Bar")
    add_body("As a shopper (Pooja),\nI want a slide-out cart drawer to open immediately when I click 'Add to Bag',\nSo that I know my item was added and can see how much more I need to spend for free shipping.")
    add_bullet("Acceptance Criteria 3.1: ", "Clicking 'Add to Bag' on any product must automatically open the slide-out Cart Drawer from the right side.")
    add_bullet("Acceptance Criteria 3.2: ", "If the subtotal is under $150, a progress bar must show the remaining amount needed for Free Shipping.")
    add_bullet("Acceptance Criteria 3.3: ", "Clickable coupon buttons (`HAUTE20`, `CONCIERGE15`) must apply the discount immediately to the subtotal.")

    # User Story 4
    add_h2("User Story 4: Single-Page Accordion Checkout")
    add_body("As a busy shopper (Genevieve),\nI want to enter my address and payment on a single clean accordion screen,\nSo that I can complete my purchase in under 60 seconds without multiple page loads.")
    add_bullet("Acceptance Criteria 4.1: ", "The checkout page (`/checkout`) must show Address, Shipping, and Payment on one accordion page.")
    add_bullet("Acceptance Criteria 4.2: ", "For testing, 1-click test card buttons must automatically fill valid card numbers (`4111 1111 1111 1111`).")
    add_bullet("Acceptance Criteria 4.3: ", "After payment, the user must immediately see an order confirmation screen with a tracking number.")

    # User Story 5
    add_h2("User Story 5: Clear, High-Contrast Ingredient Details")
    add_body("As an informed shopper (Claire),\nI want high-contrast, readable text for all ingredients and clinical results,\nSo that I can read the botanical information comfortably without eye strain.")
    add_bullet("Acceptance Criteria 5.1: ", "All body and ingredient text must have a contrast ratio of at least 4.5:1 against the background, passing WCAG 2.1 AA.")
    add_bullet("Acceptance Criteria 5.2: ", "The product details page must show a botanical active accordion with clinical efficacy percentages.")

    # Save Document
    doc.save(filename)
    os.makedirs("docs", exist_ok=True)
    doc.save(os.path.join("docs", filename))
    print(f"Strict Assignment 2 document successfully created at: {filename} and docs/{filename}")

if __name__ == "__main__":
    generate_strict_assignment2_document()
